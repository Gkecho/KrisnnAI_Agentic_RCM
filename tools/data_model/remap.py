import json, re, sys
from collections import Counter, defaultdict
from openpyxl import load_workbook

import os
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "input", "KrisnnAI-Field-Mapping.xlsx")
m = json.load(open(os.path.join(HERE, "build", "model.json"), encoding="utf-8"))
COLS = {t["name"]: {c["name"] for c in t["cols"]} for t in m["tables"]}
TBD = {t["name"] for t in m["tbd"]}

# old table -> candidate new tables, in order of preference
TM = {
    "fact_appointment": ["appointment"], "fact_visit": ["encounter"], "fact_charge_line": ["bill_details"],
    "fact_claim": ["claim"], "fact_claim_line": ["claim_details"], "fact_claim_dx": ["claim_details", "claim"],
    "fact_claim_status_event": ["claim_status_history"],
    "fact_remit_line": ["claim_details", "transactions"], "fact_remit_adjustment": ["transactions"],
    "fact_plb_adjustment": ["transactions"], "fact_posting": ["transactions"],
    "fact_payment": ["payment"], "fact_bank_deposit": ["payment"],
    "fact_patient_balance_event": ["transactions", "ar"], "fact_ar_snapshot": ["ar", "ar_interval"],
    "fact_denial_prediction": ["agent_finding", "denial"],
    "fact_agent_decision": ["agent_decision", "agent_run", "agent_decision_details"],
    "fact_finding": ["agent_finding"], "fact_work_task_event": ["agent_task"], "fact_approval": ["agent_approval"],
    "fact_writeback": ["agent_writeback"], "fact_decision_rule_result": ["agent_decision_details", "agent_decision"],
    "fact_value_ledger": ["agent_value"], "fact_kpi_value": ["agent_value"],
    "fact_auth_request_line": ["authorization"], "fact_eligibility_check": ["eligibility_check", "coverage"],
    "fact_enrollment_event": ["enrollment"], "fact_appeal_event": ["appeal"],
    "fact_document": ["document"], "fact_note": ["document"], "fact_note_element": ["document"],
    "fact_cash_forecast": ["agent_value"], "fact_coding_review_line": ["agent_finding", "encounter_diagnosis"],
    "fact_provider_credential_event": ["provider"], "fact_data_check": ["data_feed"],
    "dim_payer": ["payer"], "dim_patient": ["patient"], "dim_provider": ["provider"], "dim_coverage": ["coverage"],
    "dim_contract_rate": ["contract_rate"], "dim_adjustment_reason": ["adjustment_reason"],
    "dim_service_code": ["cpt_code"], "dim_diagnosis": ["diagnosis_code"], "dim_facility": ["facility"],
    "dim_place_of_service": ["place_of_service"], "dim_staff": ["staff"], "dim_rule": ["payer_rule"],
    "dim_metric": ["agent_value"], "dim_finding_type": ["agent_finding"], "dim_action_type": ["agent_decision"],
    "dim_decision_context": ["agent_decision"], "dim_role": ["staff"], "dim_document_type": ["document"],
    "dim_claim_status": ["claim_status_history"], "dim_ack_code": ["claim_status_history"],
    "dim_source_system": ["data_feed"], "dim_date": ["calendar"], "dim_patient_contact": ["patient"],
    "(all facts)": [],
}

# old column stem -> final column name(s)
SYN = {
    "service_code": "cpt_code", "service_code_key": "cpt_code_id", "charge_amount": "billed_amount",
    "total_charge": "billed_amount", "claim_submitter_id": "claim_number", "visit_ref": "encounter_number",
    "encounter_ref": "encounter_number", "appointment_ref": "appointment_number", "payer_claim_control_number": "payer_claim_number",
    "paid_amount": "paid_amount", "patient_responsibility": "patient_responsibility_amount",
    "diagnosis_pointer": "diagnosis_pointers", "pos": "place_of_service", "rendering_provider": "rendering_provider",
    "billing_provider": "billing_provider", "service_date": "service_date", "known_at": "known_at",
    "decision_id": "agent_decision_id", "agent_version": "agent_version", "model_id": "model_id",
    "confidence": "confidence", "is_executed": "is_executed", "subject_ref": "subject_number",
}


def norm(c):
    c = c.lower().strip()
    c = re.sub(r"_date_key$", "_date", c)
    if c.endswith("_key"):
        c = c[:-4] + "_id"
    return c


def candidates(c):
    n = norm(c)
    out = [c.lower(), n]
    base = n[:-3] if n.endswith("_id") else n
    for k in (c.lower(), n, base):
        if k in SYN:
            out.append(SYN[k])
    out += [base, base + "_id", base + "_number", base + "_code", base + "_amount"]
    return list(dict.fromkeys(out))


from overrides import OV
from overrides2 import X


def match(tbl, col):
    ov = OV.get((tbl, col))
    if ov:
        return ov[2], ov[0], ov[1], ov[3]
    x = X.get((tbl, col))
    if x:
        return x[2], x[0], x[1], x[3]
    cl = TM.get(tbl)
    if cl is None:
        return "Unmapped table", "", "", ""
    if cl and not [t for t in cl if t in COLS] and [t for t in cl if t in TBD]:
        return "Planned table", [t for t in cl if t in TBD][0], "", "table is on the TBD list, not yet designed"
    if not col:
        return ("Table only", cl[0] if cl else "", "", "")
    real = [t for t in cl if t in COLS]
    tbd = [t for t in cl if t in TBD]
    for t in real:
        if col.lower() in COLS[t]:
            return "Exact", t, col.lower(), ""
    for t in real:
        for cand in candidates(col):
            if cand in COLS[t]:
                return "Renamed", t, cand, ""
    for cand in candidates(col):
        for t, cs in COLS.items():
            if cand in cs and t not in real:
                return "Other table", t, cand, ""
    if tbd and not real:
        return "Planned table", tbd[0], "", "table is on the TBD list, not yet designed"
    if tbd:
        return "Gap", tbd[0] if not real else real[0], "", "no column; planned table " + tbd[0]
    return "Gap", real[0] if real else "", "", "no column in the final model"


if __name__ == "__main__":
    wb = load_workbook(SRC, read_only=True)
    rows = [r for r in wb["Mapping"].iter_rows(min_row=5, values_only=True) if r[0]]
    res = []
    for r in rows:
        if not r[6]:
            res.append((r, "By design", "", "", r[9] or ""))
            continue
        res.append((r,) + match(r[6], r[7]))
    c = Counter(x[1] for x in res)
    print(len(res), dict(c))
    if len(sys.argv) > 1:
        gaps = defaultdict(list)
        for r, lvl, t, col, note in res:
            if lvl in ("Gap", "Other table", "Planned table"):
                gaps[(lvl, r[6])].append(r[7])
        for k, v in sorted(gaps.items(), key=lambda kv: -len(kv[1]))[: int(sys.argv[1])]:
            print(k, len(v), sorted(set(x for x in v if x))[:12])
