import os, re
from collections import Counter, defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import remap

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.environ.get("KRISNNAI_REMAP_OUT") or os.path.join(os.path.normpath(os.path.join(HERE, "..", "..", "docs", "data_model")), "KrisnnAI-Agent-Field-Remap.xlsx")
wb = load_workbook(remap.SRC, read_only=True)
rows = [r for r in wb["Mapping"].iter_rows(min_row=5, values_only=True) if r[0]]

BYD = {"Payload", "Duplicate", "Format only", "Operational"}
res = []
for r in rows:
    old_t, old_c, treat = r[6], r[7], r[9]
    if not old_t:
        lvl, t, c, n = ("By design", "", "", f"{treat}: {r[12] or ''}".strip(": "))
    elif not old_c:
        if treat in BYD:
            lvl, n = "By design", f"{treat}: {r[12] or ''}"
        elif treat == "No home":
            lvl, n = "No home", f"Had no home in the old design either: {r[12] or ''}"
        elif treat == "Derived":
            lvl, n = "Derived", r[12] or ""
        else:
            lvl, n = "Check", r[12] or ""
        t = ""; c = ""
    else:
        lvl, t, c, n = remap.match(old_t, old_c)
    res.append((r, lvl, t, c, n))

# ---- bank deposit feed removed from Bookkeeper (A11) and Captain (A14)
REMOVE = re.compile(r"deposit_missing|duplicate_deposit|deposit_variance|deposit\.unidentified|payload: deposit_ref|a11_deposit_investigation|a11_unidentified_deposit|matched deposits|match basis|deposit lag|is_unidentified|payment-to-deposit|late_deposits|days_to_deposit|DEPOSIT_DUE|DEPOSIT_LATE|MATCH_WINDOW|UNIDENTIFIED_AFTER|per-payer deposit timing|bank_deposits\.csv|file kinds 'deposits'|bank deposit amount|scanned check|ACH trace|lockbox|funds availability|check number|remitter name|customer reference|bank reference|BAI2|type code of a transaction detail", re.I)
CHANGE = re.compile(r"gap priority|deposit_late / posting_late|rules fired|auto-post eligible|forecast cash versus|investigation outcome|cash_truth|BPR04", re.I)
res2 = []
for r, lvl, t, c, n in res:
    txt = f"{r[3]} {r[4]}"
    if r[0] in ("A11", "A14") and (r[2] == "Bank deposit feed" or r[6] == "fact_bank_deposit" or REMOVE.search(txt)) and not re.search(r"BPR|TRN|gap priority|deposit_late / posting_late", str(r[3])):
        lvl, t, c, n = "Removed", "", "", "Bank deposit feed removed from the agent"
    elif r[0] in ("A11", "A14") and CHANGE.search(txt):
        n = (n + "; " if n else "") + "REWORD: remove the bank-deposit part"
    res2.append((r, lvl, t, c, n))
res = res2
from decisions_map import DEC
res3 = []
for r, lvl, t, c, n in res:
    if lvl == "Gap" and r[7] in DEC:
        t, c, lvl, n = DEC[r[7]]
    res3.append((r, lvl, t, c, n))
res = res3
import json as _j
_PRJ = os.path.join(HERE, "build", "plan_remap.json")
PR = {}
for _k, _items in _j.load(open(_PRJ, encoding="utf-8")).items():
    for _a, _b in _items:
        PR[(_k, tuple(_a))] = _b


import json as _j2
_MJ = _j2.load(open(os.path.join(HERE, "build", "model.json"), encoding="utf-8"))
NEWT = set(_MJ["plan_tables"])


def _plan(r, lvl=''):
    t, c, ag, fl = r[6], r[7] or "", r[1], r[3]
    for kk in [("A", (t, c, ag, fl)), ("B", (t, f"{c} | {ag}: {fl}")), ("C", (t, f"{c} @ {ag} / {fl}"))]:
        if kk in PR:
            return PR[kk]
    cl = remap.TM.get(t) or []
    if (cl and set(cl) <= NEWT) or lvl == 'Planned table':
        for kk in [("C", (t, c)), ("D", (t, c)), ("B", (t, c)), ("A", (t, c))]:
            if kk in PR:
                return PR[kk]
    return None


res4 = []
nplan = 0
for r, lvl, t, c, n in res:
    if r[6] and lvl not in ("Removed", "Exact"):
        hit = _plan(r, lvl)
        if hit:
            t, c, lvl, n = hit[0], hit[1], hit[2], hit[3]
            nplan += 1
    res4.append((r, lvl, t, c, n))
res = res4
print("rows resolved from the planned-table designs:", nplan)

order = ["Exact", "Renamed", "Added", "Derived", "Partial", "Dropped", "Planned table", "Gap", "Removed", "No home", "By design", "Other table", "Check"]
MEANING = {
    "Exact": "Same column name exists in the final table.",
    "Renamed": "Final model holds it under a different name (or on the table it joins to).",
    "Derived": "Not stored; worked out from final columns.",
    "Partial": "Final model holds part of it. Note says what is missing.",
    "Planned table": "Belongs to a table on the TBD list that is not designed yet.",
    "Removed": "Dropped on purpose: the bank deposit feed is no longer an agent input.", "Added": "A new optional column was added to the model for it (see Added in review).", "Dropped": "Decided not to store it; the note says why.", "Gap": "No column holds it. Decision needed: add a column or drop the field.",
    "No home": "Carried over: the old design also had no home for it.",
    "By design": "Not a column: file, payload, format detail, duplicate or operational item.",
    "Other table": "Held on a different final table than the old design used.",
    "Check": "Needs a look.",
}
FILL = {"Exact": "E2EFDA", "Renamed": "E2EFDA", "Derived": "DDEBF7", "Partial": "FFF2CC", "Planned table": "FCE4D6",
        "Gap": "F8CBAD", "Added": "FFE699", "Dropped": "EDEDED", "Removed": "D9D9D9", "No home": "EDEDED", "By design": "EDEDED", "Other table": "FFF2CC", "Check": "F8CBAD"}
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="1F3864")


def head(ws, row, names, widths):
    for i, (h, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row, i, h); c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10); c.fill = HF
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BD
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)


def put(ws, row, vals, fill=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v); c.font = Font(name="Arial", size=9); c.border = BD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill and i == fill[0]:
            c.fill = PatternFill("solid", fgColor=fill[1])


o = Workbook(); ws = o.active; ws.title = "Read me"
cnt = Counter(x[1] for x in res)
ws["A1"] = "KrisnnAI agent field remap"; ws["A1"].font = Font(name="Arial", bold=True, size=14, color="1F3864")
ws["A2"] = "Every row of the agent field mapping (1,830) re-pointed from the old 59-table design onto the final 41-table canonical model."
info = [("Rows", len(res)), ("Distinct fields (by name)", len({r[3] for r, *_ in res})), ("", "")] + \
       [(k, f"{cnt[k]}  -  {MEANING[k]}") for k in order if cnt[k]]
for i, (a, b) in enumerate(info, 4):
    ws.cell(i, 1, a).font = Font(name="Arial", bold=True, size=10); ws.cell(i, 2, b).font = Font(name="Arial", size=10)
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 110

# Remapped sheet
ws = o.create_sheet("Remapped")
head(ws, 1, ["Area", "Agent", "Group", "Field", "What it is", "Direction", "Old table", "Old column", "Result", "Final table", "Final column", "Note"],
     [7, 16, 18, 26, 36, 10, 24, 26, 13, 24, 28, 60])
for i, (r, lvl, t, c, n) in enumerate(res, 2):
    put(ws, i, [r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], lvl, t, c, n], fill=(9, FILL[lvl]))
ws.auto_filter.ref = f"A1:L{len(res) + 1}"

# Summary by final table
ws = o.create_sheet("By final table")
head(ws, 1, ["Final table", "Rows", "Read", "Produced", "Written", "Sent", "Saved"], [28, 8, 8, 10, 9, 8, 8])
bt = defaultdict(Counter)
for r, lvl, t, c, n in res:
    if t and lvl not in ("Planned table",):
        bt[t][r[5]] += 1
        bt[t]["all"] += 1
for i, (t, cc) in enumerate(sorted(bt.items(), key=lambda kv: -kv[1]["all"]), 2):
    put(ws, i, [t, cc["all"], cc["Read"], cc["Produced"], cc["Written"], cc["Sent"], cc["Saved"]])

# Summary by agent
ws = o.create_sheet("By agent")
head(ws, 1, ["Area", "Agent", "Rows"] + order[:11], [7, 20, 8] + [11] * 11)
ba = defaultdict(Counter)
for r, lvl, t, c, n in res:
    ba[(r[0], r[1])][lvl] += 1; ba[(r[0], r[1])]["all"] += 1
for i, ((a, g), cc) in enumerate(sorted(ba.items(), key=lambda kv: (len(kv[0][0]), kv[0][0])), 2):
    put(ws, i, [a, g, cc["all"]] + [cc[k] for k in order[:11]])

# Gaps to decide
ws = o.create_sheet("Gaps to decide")
head(ws, 1, ["Final table (closest)", "Old column", "Agents using it", "Rows", "Why it matters / note", "Decision (blank)"], [26, 28, 28, 7, 80, 24])
gp = defaultdict(lambda: [set(), 0, ""])
for r, lvl, t, c, n in res:
    if lvl == "Gap":
        k = (t, r[7]); gp[k][0].add(r[1]); gp[k][1] += 1; gp[k][2] = n
for i, ((t, c), (ag, k, n)) in enumerate(sorted(gp.items(), key=lambda kv: (kv[0][0], -kv[1][1])), 2):
    put(ws, i, [t, c, ", ".join(sorted(ag)), k, n, ""])
ngap = len(gp)

# Planned tables needed
ws = o.create_sheet("Planned tables needed")
head(ws, 1, ["Planned (TBD) table", "Rows waiting", "Agents", "Old tables"], [26, 12, 50, 60])
pl = defaultdict(lambda: [0, set(), set()])
for r, lvl, t, c, n in res:
    if lvl == "Planned table":
        pl[t][0] += 1; pl[t][1].add(r[1]); pl[t][2].add(r[6])
for i, (t, (k, ag, ot)) in enumerate(sorted(pl.items(), key=lambda kv: -kv[1][0]), 2):
    put(ws, i, [t, k, ", ".join(sorted(ag)), ", ".join(sorted(x for x in ot if x))])

for w in o.worksheets:
    w.sheet_view.showGridLines = False
o.save(OUT)
print("saved", OUT)
print(dict(cnt), "| distinct gaps", ngap)
print("planned:", {t: v[0] for t, v in pl.items()})
print("check rows:", [(r[3], r[6], r[9]) for r, lvl, *_ in res if lvl == "Check"])
