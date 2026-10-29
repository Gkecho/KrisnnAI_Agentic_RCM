"""A15-lite report: one dict, written as JSON and as an Excel workbook. Money is Decimal, never float."""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from krisnnai.agents.a15_lite.detect import RECOVERABLE, STRUCTURAL, Finding
from krisnnai.agents.a15_lite.g0 import G0Result

NOTE = "Findings are detected opportunities, not recovered dollars. Nothing in this report changes any record or claim."
ZERO = Decimal(0)
OPP_COLUMNS = [
    "rank",
    "type",
    "status",
    "expected_amount",
    "agent",
    "payer",
    "claim",
    "visit",
    "line_no",
    "cpt",
    "carc",
    "category",
    "provider",
    "location",
    "deadline",
    "basis",
    "reason",
]


def _slice(findings: list[Finding], attr: str, blank: str) -> list[dict]:
    groups: dict[str, dict] = defaultdict(lambda: {"findings": 0, RECOVERABLE: ZERO, STRUCTURAL: ZERO})
    for f in findings:
        g = groups[getattr(f, attr) or blank]
        g["findings"] += 1
        g[f.status] += f.expected_amount
    rows = [
        {
            "key": k,
            "findings": g["findings"],
            RECOVERABLE: g[RECOVERABLE],
            STRUCTURAL: g[STRUCTURAL],
            "total": g[RECOVERABLE] + g[STRUCTURAL],
        }
        for k, g in groups.items()
    ]
    return sorted(rows, key=lambda r: (-r["total"], r["key"]))


def build_report(tenant_id: int, as_of: date, findings: list[Finding], g0: G0Result, kpis: dict) -> dict:
    by_type: dict[str, dict] = defaultdict(lambda: {"findings": 0, RECOVERABLE: ZERO, STRUCTURAL: ZERO})
    for f in findings:
        t = by_type[f.type]
        t["findings"] += 1
        t[f.status] += f.expected_amount
    ranked = sorted(findings, key=lambda f: (f.status != RECOVERABLE, -f.expected_amount, f.type, f.claim, f.visit))
    opportunities = [
        {"rank": i, **{k: v for k, v in f.to_dict().items() if k in OPP_COLUMNS}} for i, f in enumerate(ranked, start=1)
    ]
    return {
        "tenant_id": tenant_id,
        "as_of": as_of,
        "note": NOTE,
        "summary": {
            "findings": len(findings),
            RECOVERABLE: sum((f.expected_amount for f in findings if f.status == RECOVERABLE), ZERO),
            STRUCTURAL: sum((f.expected_amount for f in findings if f.status == STRUCTURAL), ZERO),
            "fallback_priced": sum(1 for f in findings if f.basis != "fee_schedule"),
            "by_type": dict(sorted(by_type.items())),
        },
        "by_payer": _slice(findings, "payer", "(none)"),
        "by_cpt": _slice(findings, "cpt", "(claim level)"),
        "by_location": _slice(findings, "location", "(not recorded)"),
        "by_provider": _slice(findings, "provider", "(none)"),
        "opportunities": opportunities,
        "kpis": kpis,
        "reconciliation": g0.to_dict(),
    }


def _json_default(o):
    if isinstance(o, (Decimal, date)):
        return str(o)
    raise TypeError(f"not serializable: {type(o)}")


def write_json(report: dict, path: Path) -> None:
    path.write_text(json.dumps(report, indent=2, default=_json_default), encoding="utf8")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([str(v) if isinstance(v, date) else v for v in row])
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"


def write_xlsx(report: dict, path: Path) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    s = report["summary"]
    rows = [
        ("Practice (tenant)", report["tenant_id"]),
        ("As of", str(report["as_of"])),
        ("Reconciliation (G0)", "passed" if report["reconciliation"]["passed"] else "FAILED"),
        ("Findings", s["findings"]),
        ("Recoverable $ (detected)", s[RECOVERABLE]),
        ("Structural $ (detected)", s[STRUCTURAL]),
        ("Findings priced by fallback", s["fallback_priced"]),
        ("", ""),
        ("Note", report["note"]),
    ]
    for r in rows:
        summary.append(r)
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 60
    for cell in summary["B"]:
        if isinstance(cell.value, Decimal):
            cell.number_format = "#,##0.00"
    summary.append(("", ""))
    summary.append(("Type", "Findings", "Recoverable $", "Structural $"))
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    for t, v in s["by_type"].items():
        summary.append((t, v["findings"], v[RECOVERABLE], v[STRUCTURAL]))
        for cell in summary[summary.max_row]:
            if isinstance(cell.value, Decimal):
                cell.number_format = "#,##0.00"

    slice_headers = ["", "Findings", "Recoverable $", "Structural $", "Total $"]
    for title, name, label in (
        ("By payer", "by_payer", "Payer"),
        ("By CPT", "by_cpt", "CPT"),
        ("By location", "by_location", "Location (POS)"),
        ("By provider", "by_provider", "Provider NPI"),
    ):
        _sheet(
            wb,
            title,
            [label] + slice_headers[1:],
            [[r["key"], r["findings"], r[RECOVERABLE], r[STRUCTURAL], r["total"]] for r in report[name]],
        )
    matrix = [[t, v[RECOVERABLE], v[STRUCTURAL], v[RECOVERABLE] + v[STRUCTURAL]] for t, v in s["by_type"].items()]
    _sheet(wb, "Recoverable vs structural", ["Type", "Recoverable $", "Structural $", "Total $"], matrix)
    _sheet(wb, "Opportunities", OPP_COLUMNS, [[o.get(c, "") for c in OPP_COLUMNS] for o in report["opportunities"]])
    kpi_rows = []
    for kpi, rows_ in report["kpis"].items():
        for row in rows_:
            kpi_rows.append([kpi] + [f"{k}={v}" for k, v in row.items() if v is not None])
    width = max((len(r) for r in kpi_rows), default=2)
    _sheet(wb, "Baseline KPIs", ["KPI"] + ["values"] * (width - 1), kpi_rows)
    g = report["reconciliation"]
    _sheet(
        wb,
        "Reconciliation",
        ["Check", "What it compares", "Items", "Off", "Largest difference", "Tolerance", "Passed"],
        [
            [
                c["name"],
                c["description"],
                c["items"],
                c["mismatches"],
                c["max_difference"],
                c["tolerance"],
                "yes" if c["passed"] else "NO",
            ]
            for c in g["checks"]
        ],
    )
    wb.save(path)
